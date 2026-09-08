from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "企业提交材料" / "A23_环境与部署说明_V1.0.docx"

BLUE = "17365D"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_GRAY, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None,
                 bold: bool | None = None, color: str = BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_table_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float],
              alignments: list[int] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, widths_cm)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=9, bold=True, color=WHITE)
    for ridx, values in enumerate(rows):
        row = table.add_row()
        for idx, text in enumerate(values):
            cell = row.cells[idx]
            set_cell_shading(cell, WHITE if ridx % 2 == 0 else PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = (alignments[idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(str(text))
            set_run_font(run, size=8.7, color=BLACK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.45
    if bold_lead:
        p.paragraph_format.keep_with_next = True
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.5)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.6)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10)


def add_number(doc: Document, number: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.38)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(f"{number}. {text}")
    set_run_font(run, size=10)


def remove_paragraph_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_code(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    cell = table.cell(0, 0)
    cell.width = Cm(16.2)
    set_cell_shading(cell, "F3F5F7")
    set_cell_border(cell, color="C9D1D9", size="5")
    set_cell_margins(cell, top=120, start=170, bottom=120, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=8.5, color="202124")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=8.5, color=GRAY)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def build() -> None:
    doc = Document()
    configure_styles(doc)
    doc.settings.odd_and_even_pages_header_footer = False

    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("A23 交付文档  环境与部署说明")
    set_run_font(hr, size=8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A23 蓝藻水华监测预警系统")
    set_run_font(r, size=15, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    title.add_run("环境与部署说明")
    remove_paragraph_border(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    sr = subtitle.add_run("比赛源码交付版")
    set_run_font(sr, size=12, color=GRAY)

    info = add_table(
        doc,
        ["文档项目", "内容"],
        [
            ["文档版本", "V1.0"],
            ["发布日期", "2026 年 9 月 8 日"],
            ["适用对象", "比赛评委  部署人员  项目维护人员"],
            ["适用范围", "Windows 本地部署  局域网演示  源码验收"],
        ],
        [4.2, 11.8],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    for row in info.rows[1:]:
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        row.cells[0].paragraphs[0].runs[0].bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.line_spacing = 1.45
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("本说明用于将源代码部署为可访问的网页系统。普通用户只需通过浏览器使用系统，数据采集和算法调用由程序内部完成。")
    set_run_font(r, size=10.5, color=GRAY)

    doc.add_page_break()

    add_heading(doc, "1 文档目的与部署结论")
    add_body(doc, "本系统不是单独的静态网页，而是由 Vue 前端和 FastAPI 后端共同组成。前端负责页面展示和交互，后端统一提供模拟演示接口、实时观测接口、预警接口和后续算法接口。部署时必须同时启动前端与后端。")
    add_body(doc, "普通用户无需配置或直接调用数据接口、算法接口。部署人员完成环境安装并启动服务后，用户通过浏览器访问系统即可。正式预测算法尚未接入当前业务接口，因此本版部署不把独立算法服务列为必需进程；后续接入算法后，应在本说明中补充模型文件、运行依赖、端口和健康检查。")

    add_heading(doc, "1.1 当前交付边界", 2)
    add_table(
        doc,
        ["组成", "当前情况", "部署要求"],
        [
            ["网页前端", "Vue 3 与 Vite", "必须启动或部署 dist 目录"],
            ["业务后端", "FastAPI", "必须启动  默认端口 8000"],
            ["模拟数据", "随源码提供", "用于演示轨和接口联调"],
            ["实时观测", "由 MEE 采集链生成", "在线采集或放置离线快照"],
            ["正式算法", "尚未接入业务接口", "本版不需要单独启动"],
            ["清洗数据", "正式发布物使用 Git LFS", "网页运行不要求读取全部清洗包"],
        ],
        [3.3, 5.5, 7.2],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "1.2 推荐交付结构", 2)
    add_code(doc, [
        "A23-project/",
        "  backend/                      后端服务",
        "  src/                          前端源代码",
        "  data-cleaning/                数据处理与实时采集",
        "  runtime-data/                 可选的离线运行快照",
        "  package.json                  前端依赖与脚本",
        "  README.md                     项目入口说明",
        "  一键启动系统.bat             最终交付脚本",
        "  一键停止系统.bat             最终交付脚本",
        "  环境与部署说明.docx",
    ])

    add_heading(doc, "2 环境要求")
    add_table(
        doc,
        ["类别", "最低要求", "推荐配置", "用途"],
        [
            ["操作系统", "Windows 10 64 位", "Windows 11 64 位", "比赛演示和本地部署"],
            ["Node.js", "18.0.0", "20 LTS", "前端安装与构建"],
            ["npm", "9", "随 Node LTS 安装", "前端依赖管理"],
            ["Python", "3.12", "3.12 64 位", "后端与采集程序"],
            ["浏览器", "现代浏览器", "最新版 Edge 或 Chrome", "系统访问"],
            ["内存", "8 GB", "16 GB", "安装依赖与本地运行"],
            ["磁盘", "2 GB 可用空间", "5 GB 以上", "源码 依赖 构建和日志"],
            ["网络", "离线可演示", "可访问实时公开接口", "在线更新实时观测"],
        ],
        [2.5, 3.3, 4.2, 6.0],
        [WD_ALIGN_PARAGRAPH.CENTER] * 4,
    )
    add_body(doc, "如通过 GitHub 获取包含大文件的完整源码，应安装 Git LFS 并执行 git lfs pull。网页基础运行不依赖全部原始数据；若仅进行比赛演示，可使用随包提供的运行快照。")

    add_heading(doc, "2.1 端口要求", 2)
    add_table(
        doc,
        ["服务", "默认地址", "说明"],
        [
            ["后端 API", "http://127.0.0.1:8000", "本机后端服务"],
            ["前端开发服务", "http://127.0.0.1:5173", "开发调试使用"],
            ["前端构建预览", "http://127.0.0.1:4173", "比赛演示推荐"],
            ["接口文档", "http://127.0.0.1:8000/docs", "Swagger 自动文档"],
        ],
        [3.6, 6.2, 6.2],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_page_break()

    add_heading(doc, "3 首次安装")
    add_heading(doc, "3.1 获取源码", 2)
    add_body(doc, "将源码压缩包完整解压到不需要管理员权限的目录。路径可以包含中文，但建议避免过长路径。若通过 GitHub 克隆完整仓库，先启用 Git LFS。")
    add_code(doc, [
        "git lfs install",
        "git clone <项目仓库地址>",
        "cd <项目目录>",
        "git lfs pull",
    ])

    add_heading(doc, "3.2 安装后端环境", 2)
    add_body(doc, "在项目根目录打开 PowerShell，创建项目专用 Python 虚拟环境并安装固定依赖。")
    add_code(doc, [
        "py -3.12 -m venv backend\\.venv",
        "backend\\.venv\\Scripts\\python.exe -m pip install --upgrade pip",
        "backend\\.venv\\Scripts\\python.exe -m pip install -r backend\\requirements.txt",
    ])

    add_heading(doc, "3.3 安装前端环境", 2)
    add_body(doc, "项目包含 package-lock.json，交付部署优先使用 npm ci，以便按照锁定版本安装依赖。")
    add_code(doc, [
        "npm ci",
        "npm run build",
    ])
    add_body(doc, "构建成功后，静态产物输出到 dist 目录。构建成功只代表前端资源可生成，完整验收仍需启动后端并检查接口。")

    add_heading(doc, "3.4 可选安装实时采集环境", 2)
    add_body(doc, "只有需要从公开接口在线更新实时观测时，才需要安装 data-cleaning 的完整依赖。离线演示不要求执行本步骤。")
    add_code(doc, [
        "py -3.12 -m venv data-cleaning\\.venv",
        "data-cleaning\\.venv\\Scripts\\python.exe -m pip install -r data-cleaning\\requirements.txt",
    ])

    add_heading(doc, "4 数据准备与配置")
    add_heading(doc, "4.1 数据模式", 2)
    add_table(
        doc,
        ["模式", "数据含义", "使用方式", "页面标识"],
        [
            ["在线实时观测", "启动后从公开接口采集的最新成功快照", "运行采集程序后启动后端", "实时观测并显示时间"],
            ["离线观测快照", "交付包内固定时间的历史观测快照", "复制四个目录文件后启动", "离线快照并显示时间"],
            ["模拟演示", "随源码提供的场景数据", "后端默认模拟 Provider", "SIMULATED 或模拟数据"],
        ],
        [3.1, 5.7, 4.2, 3.0],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(doc, "在线观测、离线观测快照和模拟数据必须保持明确区分。实时接口失败时，系统应显示不可用、延迟或历史快照状态，不得自动把模拟数据显示为实时观测。")

    add_heading(doc, "4.2 放置离线观测快照", 2)
    add_body(doc, "离线运行数据包应包含 observations.parquet、snapshots.json、stations.json 和 status.json。默认放置目录如下。")
    add_code(doc, [
        "data-cleaning\\storage\\silver\\mee_realtime\\observations.parquet",
        "data-cleaning\\storage\\silver\\mee_realtime\\snapshots.json",
        "data-cleaning\\storage\\silver\\mee_realtime\\stations.json",
        "data-cleaning\\storage\\silver\\mee_realtime\\status.json",
    ])
    add_body(doc, "该快照属于带时间戳的历史观测，不等同于部署当时的实时数据。交付时应同时提供数据来源、观测时间和生成时间。")

    add_heading(doc, "4.3 在线更新实时观测", 2)
    add_body(doc, "在项目根目录执行以下命令。采集成功后，程序会更新实时目录；采集失败时保留最近一次成功快照并记录失败状态。")
    add_code(doc, [
        "cd data-cleaning",
        ".venv\\Scripts\\python.exe -m data_factory collect-realtime --source mee",
        "cd ..",
    ])

    add_heading(doc, "4.4 可配置项", 2)
    add_table(
        doc,
        ["配置项", "默认值", "用途", "是否必须修改"],
        [
            ["VITE_API_BASE_URL", "/api/v1", "前端调用后端的基础路径", "本机部署无需修改"],
            ["TAIHU_REALTIME_\nCATALOG_DIR", "项目内实时目录", "指定实时目录的绝对路径", "自定义数据目录时修改"],
            ["OBSERVATION_PROVIDER", "simulated", "模拟观测 Provider 选择", "当前版本不要改为未实现值"],
            ["PREDICTION_PROVIDER", "simulated", "模拟预测 Provider 选择", "算法接入前不要修改"],
        ],
        [4.0, 3.9, 5.0, 3.1],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "5 启动与停止")
    add_heading(doc, "5.1 一键启动方式", 2)
    add_body(doc, "最终比赛交付包建议提供一键启动系统.bat。部署人员双击后，脚本应依次检查依赖和数据目录，启动后端，启动前端构建预览，执行健康检查，并自动打开浏览器。脚本不得把在线采集失败当作系统启动成功的实时数据。")
    add_number(doc, 1, "双击一键启动系统.bat。")
    add_number(doc, 2, "等待命令窗口显示后端和前端均已启动。")
    add_number(doc, 3, "确认窗口显示前端访问地址与当前数据模式。")
    add_number(doc, 4, "浏览器自动打开系统；若未自动打开，访问 http://127.0.0.1:4173。")

    add_heading(doc, "5.2 手动启动方式", 2)
    add_body(doc, "一键脚本不可用时，可在项目根目录分别打开两个 PowerShell 窗口。")
    add_body(doc, "终端一  启动后端", bold_lead="终端一")
    add_code(doc, [
        "backend\\.venv\\Scripts\\python.exe -m uvicorn backend.main:app --host 127.0.0.1 --port 8000",
    ])
    add_body(doc, "终端二  启动前端构建预览", bold_lead="终端二")
    add_code(doc, [
        "npm run build",
        "npm run preview",
    ])
    add_body(doc, "前端开发调试时可使用 npm run dev，并访问 http://127.0.0.1:5173。比赛演示优先使用构建后的预览模式，以验证实际交付产物。")

    add_heading(doc, "5.3 停止系统", 2)
    add_body(doc, "推荐使用一键停止系统.bat 关闭本项目启动的前后端进程。手动运行时，在对应终端窗口按 Ctrl C 停止服务。不要通过结束全部 Python 或 Node 进程的方式停止系统，以免影响计算机上的其他程序。")

    add_heading(doc, "6 部署验证")
    add_body(doc, "部署完成不应只检查网页是否打开。至少完成以下检查，确认前端、后端和数据链路属于同一次启动。")
    add_table(
        doc,
        ["检查项", "访问地址或操作", "通过标准"],
        [
            ["后端健康", "http://127.0.0.1:8000/api/health", "HTTP 200 且 status 为 ok"],
            ["接口文档", "http://127.0.0.1:8000/docs", "Swagger 页面可打开"],
            ["实时状态", "/api/v1/realtime/status", "明确显示 available 和 freshness_status"],
            ["实时汇总", "/api/v1/realtime/summary", "有数据时 HTTP 200  无数据时明确返回不可用"],
            ["前端首页", "http://127.0.0.1:4173", "页面资源与导航正常"],
            ["站点页面", "进入监测站点页面", "显示站点数据或明确的不可用状态"],
            ["数据身份", "检查页面标签和更新时间", "观测 模拟 离线快照不混淆"],
        ],
        [3.0, 6.3, 6.7],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "6.1 快速验收流程", 2)
    add_number(doc, 1, "启动系统并打开首页。")
    add_number(doc, 2, "进入综合驾驶舱，确认更新时间和数据模式可见。")
    add_number(doc, 3, "进入监测站点页面，选择站点并查看最新观测。")
    add_number(doc, 4, "进入历史和风险页面，确认页面能够加载或明确披露数据不可用。")
    add_number(doc, 5, "打开接口文档和实时状态接口，核对页面与接口的数据时间一致。")

    add_heading(doc, "7 常见问题")
    add_table(
        doc,
        ["现象", "可能原因", "处理方法"],
        [
            ["npm ci 失败", "Node 版本不符或网络异常", "使用 Node 20 LTS  检查网络后重试"],
            ["Python 依赖安装失败", "Python 版本或虚拟环境异常", "使用 Python 3.12 重新创建 backend .venv"],
            ["8000 端口被占用", "已有后端进程运行", "关闭旧进程或统一修改后端与代理端口"],
            ["网页打开但没有数据", "后端未启动或代理不可达", "先检查 api health 和终端错误"],
            ["实时状态为 unavailable", "没有快照或首次采集失败", "放置离线快照或运行在线采集"],
            ["实时汇总返回 409", "实时目录缺失或为空", "检查四个快照文件和目录配置"],
            ["页面刷新后接口失败", "只发布 dist 且未配置 API 代理", "使用 Vite preview 或在 Nginx IIS 中配置 api 反向代理"],
            ["局域网设备无法访问", "防火墙或访问地址错误", "开放前端端口并使用主机局域网 IP"],
            ["Git 文件下载不完整", "未安装或拉取 Git LFS", "执行 git lfs install 和 git lfs pull"],
        ],
        [4.1, 5.1, 6.8],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "8 安全与交付注意事项")
    add_bullet(doc, "不要在源码、配置示例、日志或文档中提交账号、密码、令牌和私有接口凭据。")
    add_bullet(doc, "运行数据包只保留系统演示必需的数据，不需要复制全部原始数据。")
    add_bullet(doc, "在线接口失败时应保留最近一次成功观测，并明确显示延迟、失败或不可用状态。")
    add_bullet(doc, "离线快照必须显示来源和数据时间，不得表述为部署时的实时观测。")
    add_bullet(doc, "算法接入后应补充模型文件校验、算法依赖、启动方式、接口健康检查和结果边界。")
    add_bullet(doc, "正式提交前应从干净源码目录完成一次全新安装、构建、启动和页面验收。")

    doc.add_page_break()
    add_heading(doc, "9 交付检查清单")
    add_table(
        doc,
        ["序号", "交付项目", "检查结果"],
        [
            ["1", "源代码和 package-lock.json 完整", "□"],
            ["2", "后端 requirements.txt 完整", "□"],
            ["3", "Git LFS 大文件已实际下载", "□"],
            ["4", "离线运行快照或在线采集说明已提供", "□"],
            ["5", "一键启动与停止脚本可用", "□"],
            ["6", "环境与部署说明已提供", "□"],
            ["7", "接口文档和系统使用说明已提供", "□"],
            ["8", "全新环境验收通过", "□"],
            ["9", "页面数据身份和更新时间显示正确", "□"],
            ["10", "未提交密码 密钥和个人敏感信息", "□"],
        ],
        [1.7, 11.8, 2.5],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "10 版本说明")
    add_body(doc, "本说明对应 A23 比赛源码交付版 V1.0。当前系统采用前后端双进程运行，实时观测可通过在线采集或固定快照提供；正式预测算法尚未接入业务接口。完成一键启动脚本、离线快照打包或算法接入后，应同步更新本文档版本和验证记录。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "A23 蓝藻水华监测预警系统 环境与部署说明"
    doc.core_properties.subject = "比赛源码交付版环境安装与部署说明"
    doc.core_properties.author = "A23 项目组"
    doc.core_properties.keywords = "A23 部署 环境 Vue FastAPI 实时数据"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

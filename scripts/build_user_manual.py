from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_environment_deployment_doc import (
    BLUE, GRAY, LIGHT_BLUE, add_body, add_bullet, add_code, add_heading,
    add_page_number, add_table, configure_styles, remove_paragraph_border,
    set_cell_shading, set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "企业提交材料" / "A23_系统使用说明_V1.0.docx"


def build() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("A23 交付文档  系统使用说明"), size=8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("A23 蓝藻水华监测预警系统"), size=15, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    title.add_run("系统使用说明")
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    set_run_font(subtitle.add_run("比赛演示与源码交付版"), size=12, color=GRAY)
    info = add_table(doc, ["文档项目", "内容"], [
        ["文档版本", "V1.0"],
        ["适用对象", "比赛评委  演示人员  普通使用者"],
        ["使用方式", "浏览器访问网页  无需手动调用接口"],
        ["推荐浏览器", "最新版 Edge 或 Chrome"],
        ["配套材料", "环境与部署说明  接口文档  小型运行数据包"],
    ], [4.2, 11.8], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    for row in info.rows[1:]:
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        row.cells[0].paragraphs[0].runs[0].bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("系统启动后，普通使用者只需通过左侧导航进入页面并查看数据，不需要配置算法或调用后端接口。"), size=10.5, color=GRAY)

    doc.add_page_break()
    add_heading(doc, "1 快速开始")
    add_body(doc, "部署人员完成启动后，使用者打开系统地址即可进入首页。本地开发模式通常访问 http://127.0.0.1:5173；构建预览模式通常访问 http://127.0.0.1:4173。企业或比赛现场部署时，以部署人员提供的实际地址为准。")
    add_table(doc, ["步骤", "操作", "预期结果"], [
        ["1", "打开系统网址", "显示系统首页和数据身份信息"],
        ["2", "查看实时状态", "确认数据可用性、观测时间和新鲜度"],
        ["3", "从左侧导航选择页面", "进入驾驶舱、站点、推演、历史、大屏或预警页面"],
        ["4", "使用页面筛选或时间轴", "查看指定站点、年份或快照"],
        ["5", "遇到错误时点击重试", "重新请求后端，不把错误理解为零值"],
    ], [2.0, 6.2, 7.8], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "1.1 演示人员启动方式", 2)
    add_body(doc, "本节供演示人员使用，普通浏览用户可以跳过。先启动后端，再启动前端；两个窗口均保持运行。")
    add_code(doc, [
        "# 窗口一  后端",
        "backend\\.venv\\Scripts\\python.exe -m uvicorn backend.main:app --host 127.0.0.1 --port 8000",
        "",
        "# 窗口二  前端",
        "npm run dev",
        "",
        "# 浏览器访问 http://127.0.0.1:5173",
    ])
    add_body(doc, "如果使用随源码提供的小型运行数据包，应在启动后端前按《小型运行数据包与数据接入说明》设置 TAIHU_REALTIME_CATALOG_DIR。")

    add_heading(doc, "2 页面导航")
    add_table(doc, ["导航", "主要内容", "典型操作"], [
        ["首页", "系统入口和实时筛查摘要", "点击业务入口或预警站点"],
        ["综合驾驶舱", "全湖地图、健康度、预警和关键指标", "选站、播放快照、打开站点详情"],
        ["监测站点", "逐站观测、质量、趋势和事件", "选择站点、查看指标和质量状态"],
        ["时空推演", "遥感影像、年份对比和规则研判", "切图层、选年份、拖动时间轴"],
        ["历史复盘", "实时快照列表和明细", "设置日期、选择快照、查看站点观测"],
        ["实时大屏", "站点状态墙和汇总指标", "浏览异常、缺测和预警站点"],
        ["预警通知", "当前预警、历史和推送状态", "筛选预警；授权人员可手动评估"],
    ], [3.2, 7.1, 5.7], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "3 综合驾驶舱")
    add_body(doc, "综合驾驶舱是默认的全湖研判页面。左侧地图显示站点位置和风险颜色，右侧卡片显示湖体健康、蓝藻筛查预警和关键指标；底部时间轴用于回放已经发布的实时快照。")
    add_bullet(doc, "点击地图站点或预警列表中的站点，可聚焦该站并打开详情。")
    add_bullet(doc, "点击时间轴刻度可查看指定快照；点击播放按钮可按快照顺序回放。")
    add_bullet(doc, "历史快照必须按其观测时间理解，不能当作当前实时状态。")
    add_bullet(doc, "颜色表示页面定义的风险或状态等级；应同时查看数据身份、观测时间和质量信息。")

    add_heading(doc, "4 监测站点")
    add_body(doc, "监测站点页面用于查看单个 MEE 观测站点。先从站点列表或地图选择站点，再查看最新值、指标趋势、缺测原因、质量状态和快照积累情况。")
    add_table(doc, ["区域", "怎么看", "注意事项"], [
        ["站点列表", "选择需要研判的站点", "站点集合以最新成功快照为准"],
        ["实时地图", "查看站点位置和风险着色", "坐标状态应结合质量信息判断"],
        ["观测指标", "查看数值、单位和观测时间", "null 表示缺测，不等于 0"],
        ["趋势面板", "查看同一指标的时间变化", "样本不足时只显示点，不做插值"],
        ["质量面板", "查看覆盖率、缺测和 suitability", "接口可用不代表数据适合建模"],
        ["事件与预警", "查看事件或执行情景流程", "情景操作不等于正式政府预警"],
    ], [3.7, 6.3, 6.0], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "5 时空推演")
    add_body(doc, "时空推演页面同时承载年度遥感影像和站点时空信息。进入页面后先选择模式，再选择图层、年份或时间节点。")
    add_bullet(doc, "遥感模式可在叶绿素 a 和漂浮藻类图层之间切换，并选择影像年份。")
    add_bullet(doc, "开启年份对比后，分别选择影像 A 和影像 B，页面显示对比摘要。")
    add_bullet(doc, "时间轴中的历史节点对应观测快照；未来节点是规则研判或模拟能力，不能表述为已上线正式预测模型。")
    add_bullet(doc, "当前风险面可能为空，系统不会为展示效果虚构湖岸边界、面积或迁移路径。")

    add_heading(doc, "6 历史复盘")
    add_body(doc, "历史复盘页面用于按日期查看已积累的实时快照。设置开始日期和结束日期后点击应用，从左侧快照列表选择记录，右侧显示快照摘要、筛查预警和逐站观测。")
    add_bullet(doc, "日期范围最大为 90 天；超过范围会提示错误。")
    add_bullet(doc, "点击重置可清除筛选；选择快照后可继续前后切换。")
    add_bullet(doc, "若只有一个快照，页面可以显示当前记录，但无法形成连续历史趋势。")

    add_heading(doc, "7 实时大屏")
    add_body(doc, "实时大屏适合比赛现场展示。页面将活跃站点、观测时间、缺测率、预警数量和逐站状态集中显示。使用者主要负责观察，不需要逐一操作。")
    add_bullet(doc, "出现 delayed 或 severely overdue 时，说明数据已经延迟或严重过期，不应继续称为当前实时状态。")
    add_bullet(doc, "站点缺测应保留缺测标识，不要用零值或模拟值替代。")
    add_bullet(doc, "大屏与驾驶舱、站点页共用实时数据链路；若多页同时异常，应优先检查后端和实时目录。")

    add_heading(doc, "8 预警通知")
    add_body(doc, "右上角铃铛和预警通知页显示当前生效预警、历史记录、渠道状态和推送回执。页面中的阈值和等级用于当前系统筛查口径。")
    add_bullet(doc, "普通使用者可以查看和筛选预警，不需要手动执行评估。")
    add_bullet(doc, "点击立即评估会触发后端重新评估并写入预警状态。若部署人员已启用邮件或短信通道，还可能产生对外投递，因此只允许授权维护人员操作。")
    add_bullet(doc, "情景预警、真实观测筛查和正式监管预警必须分开描述。比赛演示结果不等同于监管部门发布。")

    add_heading(doc, "9 数据状态怎么看")
    add_table(doc, ["标识", "含义", "使用建议"], [
        ["observed", "来自实时观测发布目录", "结合 as_of、freshness 和质量状态使用"],
        ["simulated", "固定演示或规则生成数据", "只用于页面演示和接口联调"],
        ["normal", "观测滞后不超过 6 小时", "可作为当前页面展示"],
        ["delayed", "观测滞后超过 6 小时且不超过 12 小时", "明确显示延迟"],
        ["severely_overdue", "观测滞后超过 12 小时", "不得称为当前实时状态"],
        ["unavailable", "没有可读取的成功快照", "检查数据目录和采集状态"],
        ["missing", "某项观测缺失", "保留缺测原因，不用 0 替代"],
    ], [4.0, 7.0, 5.0], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "10 常见问题")
    add_table(doc, ["现象", "可能原因", "处理方法"], [
        ["页面打不开", "前端未启动或地址错误", "确认前端窗口运行并使用正确端口"],
        ["页面能打开但没有业务数据", "后端未启动或代理失败", "检查 8000 端口和 /api/health"],
        ["实时数据不可用", "目录缺失、为空或发布不一致", "检查 /realtime/status 和四个目录文件"],
        ["数据时间较旧", "上游暂未更新或采集失败", "查看 freshness、last_success_at 和 last_error"],
        ["预测按钮不可用", "当前能力受限或正式算法未接入", "以能力接口和页面提示为准，不强行演示"],
        ["遥感图层不显示", "图层文件或 manifest 不可读", "点击重试并检查 rs_overlays"],
        ["日期筛选报错", "日期顺序错误或超过 90 天", "调整起止日期后重新应用"],
    ], [4.2, 5.8, 6.0], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "11 演示建议")
    add_body(doc, "推荐演示顺序为：首页确认系统和数据身份，进入综合驾驶舱展示全湖态势，选择一个站点查看观测与质量，进入历史复盘展示快照，再进入时空推演展示遥感年份对比，最后打开预警通知说明状态和边界。")
    add_bullet(doc, "演示前先检查后端健康、实时状态和浏览器窗口尺寸。")
    add_bullet(doc, "现场只使用已经验证过的快照和页面，不临时更换数据目录。")
    add_bullet(doc, "算法尚未正式接入时，应表述为接口和页面已经预留，不表述为模型已经上线。")
    add_bullet(doc, "出现数据缺失、延迟或能力不可用时，直接展示系统的明确提示，这本身也是可信交付的一部分。")

    add_heading(doc, "12 版本说明")
    add_body(doc, "本说明对应当前 Vue 前端、FastAPI OpenAPI 2.1.0 和比赛交付数据包。后续页面名称、数据字段、预警通道或正式算法能力发生变化时，应同步更新本说明。")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

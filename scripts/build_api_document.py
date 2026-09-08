from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_environment_deployment_doc import (
    BLUE,
    GRAY,
    LIGHT_BLUE,
    add_body,
    add_bullet,
    add_code,
    add_heading,
    add_page_number,
    add_table,
    configure_styles,
    remove_paragraph_border,
    set_cell_shading,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "企业提交材料" / "A23_接口文档_V1.0.docx"


def add_endpoint_heading(doc: Document, method: str, path: str, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{method}  {path}")
    set_run_font(r, name="Consolas", size=10.5, bold=True, color=BLUE)
    r2 = p.add_run(f"  {title}")
    set_run_font(r2, size=10.5, bold=True)


def prevent_row_split(table) -> None:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def add_param_table(doc: Document, rows: list[list[str]]) -> None:
    table = add_table(
        doc,
        ["参数", "位置", "必填", "类型与取值", "说明"],
        rows,
        [3.0, 1.8, 1.5, 4.2, 5.5],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    prevent_row_split(table)


def add_field_table(doc: Document, rows: list[list[str]]) -> None:
    table = add_table(
        doc,
        ["字段", "类型", "说明"],
        rows,
        [4.4, 2.8, 8.8],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    prevent_row_split(table)


def build() -> None:
    doc = Document()
    configure_styles(doc)
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("A23 交付文档  接口文档")
    set_run_font(run, size=8.5, color="000000")
    add_page_number(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A23 蓝藻水华监测预警系统")
    set_run_font(r, size=15, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    title.add_run("接口文档")
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(46)
    r = subtitle.add_run("比赛源码交付版")
    set_run_font(r, size=12, color=GRAY)

    cover_table = add_table(
        doc,
        ["文档项目", "内容"],
        [
            ["文档版本", "V1.0"],
            ["API 版本", "2.1.0"],
            ["基础路径", "/api/v1"],
            ["发布日期", "2026 年 9 月 8 日"],
            ["适用对象", "前端开发  后端维护  部署与评审人员"],
        ],
        [4.2, 11.8],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    for row in cover_table.rows[1:]:
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        row.cells[0].paragraphs[0].runs[0].bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.line_spacing = 1.45
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("本文档描述当前源码版的实际接口契约。普通用户无需手动调用接口，网页由前端程序自动调用。")
    set_run_font(r, size=10.5, color=GRAY)

    doc.add_page_break()

    add_heading(doc, "1 接口概述")
    add_body(doc, "系统后端使用 FastAPI 对前端提供 JSON 接口。本地直连地址为 http://127.0.0.1:8000，业务接口基础路径为 /api/v1。前端默认使用相对路径 /api/v1，由 Vite 或正式 Web 服务器代理到后端。")
    add_body(doc, "当前比赛版未实现登录、令牌和角色权限，仅适合本机或可信局域网演示。如对外网发布，应先增加 HTTPS、身份认证、权限控制、请求限流和日志脱敏。")

    add_heading(doc, "1.1 调用信息", 2)
    add_table(
        doc,
        ["项目", "约定"],
        [
            ["直连根地址", "http://127.0.0.1:8000"],
            ["业务基础路径", "/api/v1"],
            ["前端默认配置", "VITE_API_BASE_URL=/api/v1"],
            ["接口文档", "http://127.0.0.1:8000/docs"],
            ["数据格式", "application/json  UTF-8"],
            ["时间格式", "ISO 8601  可包含时区"],
            ["当前认证", "无  仅用于本地和可信网络"],
        ],
        [5.0, 11.0],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "1.2 能力边界", 2)
    add_table(
        doc,
        ["数据轨", "版本", "当前能力", "不得宣称"],
        [
            ["实时观测", "MEE-RT-V1", "MEE 国控站点观测、汇总、质量和时间轴", "未经跨源验证的科学真值"],
            ["模拟演示", "DEMO-OBS-V1", "演示分区、观测和页面联调", "真实站点或监管数据"],
            ["模拟预测", "DEMO-PRED-V1", "T+1 3 7 15 天样例结果", "已上线的正式预测模型"],
            ["遥感图层", "THQBCA-V2-BIOOPTICS-V1", "历史年度遥感反演图层", "实时卫星观测或地面实测"],
        ],
        [3.0, 4.0, 5.6, 3.4],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "2 通用响应契约")
    add_body(doc, "除 Swagger 静态资源外，业务接口统一返回信封结构。成功和失败响应都保留数据模式、数据版本、时间和请求追踪信息。")
    add_code(doc, [
        "{",
        '  "code": 200,',
        '  "message": "ok",',
        '  "data": { ... },',
        '  "meta": {',
        '    "data_mode": "observed",',
        '    "dataset_version": "MEE-RT-V1",',
        '    "prediction_run_id": null,',
        '    "as_of": "2026-09-08T02:24:34Z",',
        '    "claim_boundary": "official_observation_not_cross_validated",',
        '    "request_id": "req_xxx"',
        "  },",
        '  "errors": []',
        "}",
    ])
    add_field_table(doc, [
        ["code", "integer", "HTTP 业务状态码  成功通常为 200"],
        ["message", "string", "响应摘要或错误提示"],
        ["data", "object array null", "业务数据  错误时通常为 null"],
        ["meta.data_mode", "string", "observed simulated 或其他明确模式"],
        ["meta.dataset_version", "string", "当前响应所属数据集版本"],
        ["meta.prediction_run_id", "string null", "预测或演示运行标识"],
        ["meta.as_of", "string", "数据截止或观测时间"],
        ["meta.claim_boundary", "string", "数据或模型能力的使用边界"],
        ["meta.request_id", "string", "单次请求追踪标识"],
        ["errors", "array", "错误列表  成功时为空数组"],
    ])

    add_heading(doc, "2.1 错误响应示例", 2)
    add_code(doc, [
        "{",
        '  "code": 409,',
        '  "message": "实时观测数据不可用",',
        '  "data": null,',
        '  "meta": { "data_mode": "observed", "dataset_version": "MEE-RT-V1", ... },',
        '  "errors": [{',
        '    "code": "REALTIME_DATA_UNAVAILABLE",',
        '    "field": null,',
        '    "detail": "实时目录不存在或为空"',
        "  }]",
        "}",
    ])

    add_heading(doc, "3 接口总览")
    add_body(doc, "当前 OpenAPI 定义共包含 32 个操作。下表为完整路由清单，后续章节对前端联调中最重要的接口进行详细说明。")
    overview_rows = [
        ["GET", "/", "后端根状态", "系统"],
        ["GET", "/api/health", "健康检查", "系统"],
        ["GET", "/api/v1/system/capabilities", "能力与阻断项", "系统"],
        ["GET", "/api/v1/datasets/summary", "数据集摘要", "模拟"],
        ["GET", "/api/v1/pipeline/runs/latest", "最新流水线运行", "模拟"],
        ["GET", "/api/v1/dashboard/overview", "页面总览", "模拟"],
        ["GET", "/api/v1/spatial-entities", "站点或演示分区列表", "双轨"],
        ["GET", "/api/v1/realtime/status", "实时链路状态", "观测"],
        ["GET", "/api/v1/realtime/summary", "全湖实时汇总或回放", "观测"],
        ["GET", "/api/v1/realtime/timeline", "实时快照时间轴", "观测"],
        ["GET", "/api/v1/realtime/alerts/overview", "预警总览", "观测"],
        ["GET", "/api/v1/realtime/alerts", "预警历史", "观测"],
        ["POST", "/api/v1/realtime/alerts/evaluate", "手动触发预警评估", "观测 写"],
        ["GET", "/api/v1/spatial-entities/{entity_id}", "空间对象详情", "双轨"],
        ["GET", "/api/v1/spatial-entities/{entity_id}/observations", "观测值", "双轨"],
        ["GET", "/api/v1/spatial-entities/{entity_id}/quality", "数据质量", "双轨"],
        ["GET", "/api/v1/forecast-capabilities", "预测能力状态", "模拟"],
        ["GET", "/api/v1/forecasts", "查询预测", "模拟"],
        ["GET", "/api/v1/forecasts/{forecast_id}", "预测详情", "模拟"],
        ["GET", "/api/v1/forecasts/{forecast_id}/explanations", "预测解释", "模拟"],
        ["GET", "/api/v1/map/layers", "地图图层清单", "混合元数据"],
        ["GET", "/api/v1/rs/manifest", "遥感图层清单", "历史遥感"],
        ["GET", "/api/v1/map/risk-polygons", "风险面要素", "模拟 当前为空"],
        ["GET", "/api/v1/events", "规范事件列表", "模拟"],
        ["GET", "/api/v1/cockpit/time-stages", "驾驶舱时间档位", "模拟"],
        ["GET", "/api/v1/cockpit/points", "驾驶舱点位", "模拟"],
        ["GET", "/api/v1/cockpit/points/{entity_id}", "驾驶舱点位详情", "模拟"],
        ["GET", "/api/v1/cockpit/risk-heatmap", "演示热力场", "模拟"],
        ["GET", "/api/v1/cockpit/events", "驾驶舱兼容事件", "模拟"],
        ["GET", "/api/v1/cockpit/region-summary", "分区摘要", "模拟"],
        ["POST", "/api/v1/cockpit/handle-warning", "处理演示事件", "模拟 写"],
        ["GET", "/api/v1/cockpit/timeline", "驾驶舱时间线", "模拟"],
    ]
    overview = add_table(
        doc,
        ["方法", "路径", "用途", "轨道"],
        overview_rows,
        [1.5, 7.1, 4.8, 2.6],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    prevent_row_split(overview)

    doc.add_page_break()
    add_heading(doc, "4 实时观测接口")
    add_body(doc, "实时观测接口读取 MEE 快照目录。当目录缺失或从未成功采集时，/status 仍返回 200 并显示 available=false；需要业务数据的其他实时接口返回 409 REALTIME_DATA_UNAVAILABLE。")

    add_endpoint_heading(doc, "GET", "/api/v1/realtime/status", "查询采集与数据新鲜度")
    add_body(doc, "无请求参数。建议页面状态栏和启动检查首先调用该接口。")
    add_field_table(doc, [
        ["available", "boolean", "当前是否有可读取的成功快照"],
        ["collection_status", "string", "最近采集状态"],
        ["freshness_status", "string", "normal delayed stale unavailable 等新鲜度状态"],
        ["snapshot_count", "integer", "已发布快照数"],
        ["active_station_count", "integer", "最新快照中活跃站点数"],
        ["latest_observed_at", "string null", "最新观测时间"],
        ["observed_lag_h", "number null", "观测滞后小时数"],
        ["latest_snapshot_id", "string null", "最新成功快照标识"],
        ["last_error", "string null", "最近一次采集错误"],
    ])

    add_endpoint_heading(doc, "GET", "/api/v1/realtime/summary", "全湖实时汇总")
    add_param_table(doc, [["snapshot", "query", "否", "string", "指定历史快照 ID  缺省为最新快照"]])
    add_body(doc, "返回站点总数、水质类别分布、达标率、蓝藻筛查预警、均值、趋势、健康分和地图标记。快照回放不得显示为当前实时状态。")
    add_code(doc, [
        "GET /api/v1/realtime/summary",
        "GET /api/v1/realtime/summary?snapshot=mee_surface_water_realtime_YYYYMMDDThhmmssZ",
    ])

    add_endpoint_heading(doc, "GET", "/api/v1/realtime/timeline", "快照时间轴")
    add_body(doc, "无请求参数。返回 snapshots 数组、latest_snapshot_id 和预警阈值口径，用于驾驶舱和历史页面回放。")

    add_endpoint_heading(doc, "GET", "/api/v1/spatial-entities", "实时站点或模拟分区列表")
    add_param_table(doc, [
        ["mode", "query", "否", "observed simulated  默认 simulated", "选择观测站点或演示分区"],
        ["entity_type", "query", "否", "string", "observed 时应为 monitoring_station"],
        ["active", "query", "否", "latest all  默认 latest", "站点是否只返回最新快照活跃集"],
        ["province", "query", "否", "string", "按省级名称筛选"],
        ["location_status", "query", "否", "verified metadata_only suspicious missing", "按坐标可信状态筛选"],
    ])
    add_code(doc, [
        "GET /api/v1/spatial-entities?mode=observed&active=latest",
        "GET /api/v1/spatial-entities?mode=observed&province=江苏省&location_status=verified",
    ])

    add_endpoint_heading(doc, "GET", "/api/v1/spatial-entities/{entity_id}/observations", "站点观测值")
    add_param_table(doc, [
        ["entity_id", "path", "是", "string", "实时站点通常为 mee-xxxxxxxx  模拟分区为分区 ID"],
        ["window", "query", "否", "latest range  默认 latest", "最新值或时间范围"],
        ["start", "query", "range 时是", "ISO 日期", "范围起日"],
        ["end", "query", "range 时是", "ISO 日期", "范围止日  最大跨度 90 天"],
        ["variables", "query", "否", "逗号分隔字符串", "实时轨需返回的指标代码"],
        ["variable_code", "query", "否", "string", "模拟轨单指标筛选"],
    ])
    add_field_table(doc, [
        ["observed_at", "string", "上游观测时间"],
        ["retrieved_at", "string", "本系统抓取时间"],
        ["variable_code", "string", "规范指标代码"],
        ["value", "number null", "观测值  缺测时为 null"],
        ["unit", "string", "规范单位"],
        ["observation_status", "string", "ok missing qc_rejected parse_failed 等"],
        ["missing_reason", "string null", "缺测原因  例如 upstream_missing"],
        ["qc_status", "string", "范围质控状态"],
        ["is_ground_truth", "boolean", "当前官方接口观测仍为 false"],
    ])

    add_endpoint_heading(doc, "GET", "/api/v1/spatial-entities/{entity_id}/quality", "站点数据质量")
    add_body(doc, "返回指标覆盖率、缺测指标、质控拒绝项、滞后时间、坐标可信状态和 suitability。suitability.model_use=false 表示当前数据未被认定为模型真值输入。")

    add_heading(doc, "5 预警接口")
    add_endpoint_heading(doc, "GET", "/api/v1/realtime/alerts/overview", "预警总览")
    add_body(doc, "返回预警功能是否启用、通道状态、生效告警数、最近告警、最后评估时间和阈值口径。")

    add_endpoint_heading(doc, "GET", "/api/v1/realtime/alerts", "预警事件列表")
    add_param_table(doc, [
        ["status", "query", "否", "active resolved all  默认 all", "告警状态"],
        ["level", "query", "否", "light moderate", "告警级别"],
        ["limit", "query", "否", "1 至 1000  默认 200", "返回条数上限"],
    ])

    add_endpoint_heading(doc, "POST", "/api/v1/realtime/alerts/evaluate", "手动触发预警巡检")
    add_body(doc, "无请求体。该接口会执行一次评估并写入告警状态。如 backend/alerts-config.json 启用了真实邮件或短信通道，还可能产生对外投递，仅限授权维护人员调用。")

    add_heading(doc, "6 预测与解释接口")
    add_body(doc, "当前预测 Provider 为 simulated。返回结果用于页面联调和比赛演示，不代表正式算法已接入，也不得用于真实预警发布。")

    add_endpoint_heading(doc, "GET", "/api/v1/forecast-capabilities", "预测能力表")
    add_body(doc, "前端应优先读取该接口决定按钮和页面文案，而不应仅根据 HTTP 200 推断某项能力已可用。")

    add_endpoint_heading(doc, "GET", "/api/v1/forecasts", "查询模拟预测")
    add_param_table(doc, [
        ["spatial_entity_id", "query", "是", "string", "已注册的 demo_zone ID"],
        ["horizon_days", "query", "否", "1 3 7 15 30  默认 3", "30 天当前返回能力阻塞"],
        ["target_metric", "query", "否", "string  默认 bloom_risk", "目标指标字段  当前 Provider 固定生成蓝藻风险"],
    ])
    add_field_table(doc, [
        ["id", "string", "预测稳定 ID"],
        ["spatial_entity_id", "string", "演示分区 ID"],
        ["horizon_days", "integer", "预测时效"],
        ["risk_score", "number", "模拟风险分"],
        ["risk_level", "string", "风险级别"],
        ["provider_type", "string", "当前为 simulation"],
        ["model_version", "string", "当前为演示规则版本"],
        ["uncertainty", "object", "演示不确定区间与方法"],
        ["quality_gate", "object", "质量门禁决策和原因"],
    ])
    add_code(doc, [
        "GET /api/v1/forecasts?spatial_entity_id=northwest_hotspot&horizon_days=3",
        "GET /api/v1/forecasts/demo-forecast-northwest_hotspot-3d",
        "GET /api/v1/forecasts/demo-forecast-northwest_hotspot-3d/explanations",
    ])
    add_body(doc, "解释接口当前返回 demo_rule_contribution，不是 SHAP。算法正式接入后，应保留路由不变，替换 Provider 实现并更新数据版本、模型版本和能力边界。")

    add_heading(doc, "7 地图与遥感接口")
    add_endpoint_heading(doc, "GET", "/api/v1/map/layers", "图层清单")
    add_body(doc, "返回年度叶绿素 a 遥感反演、漂浮藻类覆盖率反演和 MEE 实时站点图层元数据。operational_use=false 表示当前不是正式生产运行图层。")

    add_endpoint_heading(doc, "GET", "/api/v1/rs/manifest", "遥感图层清单")
    add_body(doc, "返回 generated_at、source_dataset、note 和 layers。如 backend/rs_overlays/manifest.json 不存在或不可读，返回 409 CAPABILITY_UNAVAILABLE。")

    add_endpoint_heading(doc, "GET", "/api/v1/map/risk-polygons", "风险面要素")
    add_param_table(doc, [["horizon_days", "query", "否", "1 3 7 15 30  默认 3", "风险时效"]])
    add_body(doc, "当前诚实返回空 FeatureCollection，不生成虚构湖岸边界、面积或迁移路径。")

    add_heading(doc, "8 驾驶舱兼容接口")
    add_body(doc, "cockpit 路由为现有驾驶舱页面提供兼容数据，当前均属于模拟演示轨。新的实时站点页应优先使用 realtime 和 spatial-entities 接口。")
    add_table(
        doc,
        ["路径", "关键参数", "说明"],
        [
            ["/cockpit/time-stages", "无", "驾驶舱时间档位"],
            ["/cockpit/points", "无", "分区点位和演示指标"],
            ["/cockpit/points/{entity_id}", "entity_id", "指定分区点位详情"],
            ["/cockpit/risk-heatmap", "无", "演示热力场"],
            ["/cockpit/events", "无", "驾驶舱事件兼容视图"],
            ["/cockpit/region-summary", "无", "分区汇总"],
            ["/cockpit/timeline", "start end", "ISO 日期  最大跨度 90 天"],
            ["/cockpit/handle-warning", "JSON event_id", "处理演示事件  POST"],
        ],
        [6.2, 3.4, 6.4],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_code(doc, [
        "POST /api/v1/cockpit/handle-warning",
        "Content-Type: application/json",
        "",
        '{ "event_id": "demo-event-1" }',
    ])

    add_heading(doc, "9 错误码")
    error_table = add_table(
        doc,
        ["HTTP", "错误码", "含义", "建议处理"],
        [
            ["404", "ENTITY_NOT_FOUND", "对象或预测记录不存在", "刷新列表并检查 ID"],
            ["409", "FORECAST_NOT_AVAILABLE", "当前 Provider 无法生成记录", "显示能力不可用"],
            ["409", "CAPABILITY_UNAVAILABLE", "能力未就绪", "根据 detail 隐藏或禁用功能"],
            ["409", "SIMULATION_ONLY", "请求真实能力但当前仅有模拟轨", "保留模拟标识"],
            ["409", "REALTIME_DATA_UNAVAILABLE", "实时快照未发布", "显示不可用并检查 status"],
            ["422", "INVALID_DATE_RANGE", "日期缺失、无法解析或顺序错误", "更正 start 和 end"],
            ["422", "QUERY_RANGE_TOO_LARGE", "日期跨度超过 90 天", "拆分查询"],
            ["422", "INVALID_HORIZON", "预测时效不支持", "使用 1 3 7 15 30"],
            ["422", "INVALID_EVENT_ID", "演示事件引用不存在", "从事件列表重新取 ID"],
            ["422", "REQUEST_VALIDATION_FAILED", "参数类型或格式不合法", "检查 errors.field"],
            ["500", "INTERNAL_ERROR", "未捕获服务器错误", "记录 request_id 并检查后端日志"],
        ],
        [1.4, 4.6, 5.4, 4.6],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    prevent_row_split(error_table)

    add_heading(doc, "10 前端接入示例")
    add_body(doc, "前端使用统一的请求层访问后端。下列示例只展示基本错误处理，实际项目应继续使用 src/services/api.js 和 src/services/realtime.js，避免页面各自重复实现信封解析。")
    add_code(doc, [
        "async function getRealtimeStatus() {",
        "  const response = await fetch('/api/v1/realtime/status')",
        "  const payload = await response.json()",
        "  if (!response.ok) {",
        "    const error = payload.errors?.[0]",
        "    throw new Error(error?.code || payload.message)",
        "  }",
        "  return payload.data",
        "}",
    ])

    add_heading(doc, "10.1 接入规则", 2)
    add_bullet(doc, "使用 meta.data_mode 和 meta.claim_boundary 决定页面数据标识。")
    add_bullet(doc, "不得在实时请求失败时静默切换为模拟数据。")
    add_bullet(doc, "使用 meta.request_id 辅助日志追踪，不将完整后端错误堆栈展示给用户。")
    add_bullet(doc, "调用写接口前显示用户确认，并防止连续重复提交。")
    add_bullet(doc, "接口可用不等于数据适合建模，仍需读取 quality 和 suitability。")

    add_heading(doc, "11 联调检查清单")
    add_table(
        doc,
        ["序号", "检查项", "通过标准"],
        [
            ["1", "后端健康", "/api/health 返回 HTTP 200"],
            ["2", "Swagger", "/docs 可打开并显示 API 2.1.0"],
            ["3", "统一信封", "code message data meta errors 齐全"],
            ["4", "实时状态", "status 如实返回 available 和 freshness_status"],
            ["5", "无数据处理", "实时汇总返回明确 409  不回退模拟"],
            ["6", "数据身份", "observed simulated 遥感历史不混淆"],
            ["7", "日期门禁", "错误日期和超 90 天返回 422"],
            ["8", "预测边界", "T+30 返回能力阻塞  演示预测保留 simulation_only"],
            ["9", "写接口", "评估和处理事件均有用户确认"],
            ["10", "追踪", "问题记录 request_id 与请求时间"],
        ],
        [1.5, 5.2, 9.3],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "12 版本说明")
    add_body(doc, "本文档对应后端 OpenAPI 2.1.0 和当前源码实现。当接入正式算法、增加登录权限、调整路由、变更响应字段或开放对外调用时，应同步更新 API 版本、OpenAPI 定义和本文档。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "A23 蓝藻水华监测预警系统 接口文档"
    doc.core_properties.subject = "比赛源码交付版 API 契约"
    doc.core_properties.author = "A23 项目组"
    doc.core_properties.keywords = "A23 API FastAPI 实时观测 模拟预测"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

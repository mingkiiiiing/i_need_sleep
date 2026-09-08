from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_environment_deployment_doc import (
    BLUE, GRAY, LIGHT_BLUE, add_body, add_bullet, add_code, add_heading,
    add_page_number, add_table, configure_styles, remove_paragraph_border,
    set_cell_shading, set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "企业提交材料" / "A23_小型运行数据包与数据接入说明_V1.0.docx"


def build() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("A23 交付文档  数据包与接入说明"), size=8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("A23 蓝藻水华监测预警系统"), size=15, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    title.add_run("小型运行数据包与数据接入说明")
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    set_run_font(subtitle.add_run("比赛源码交付版"), size=12, color=GRAY)
    info = add_table(doc, ["文档项目", "内容"], [
        ["文档版本", "V1.0"],
        ["数据包名称", "A23 小型运行数据包 V1.0"],
        ["用途", "离线运行  比赛演示  接口联调"],
        ["实时数据版本", "MEE-RT-V1"],
        ["模拟数据版本", "DEMO-OBS-V1  DEMO-PRED-V1"],
    ], [4.2, 11.8], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    for row in info.rows[1:]:
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        row.cells[0].paragraphs[0].runs[0].bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("本说明回答两个问题：评审如何拿到一份可运行的数据，以及后续真实数据应按什么目录和字段接入。"), size=10.5, color=GRAY)

    doc.add_page_break()
    add_heading(doc, "1 数据包内容")
    add_body(doc, "数据包已经放在企业提交材料目录中，解压后不需要改动源代码即可使用。实时目录提供当前后端真正读取的四件套；模拟目录只提供字段示例，不会自动覆盖实时数据。")
    add_table(doc, ["目录或文件", "作用", "是否进入实时链路"], [
        ["realtime_catalog/stations.json", "站点和空间对象目录", "是"],
        ["realtime_catalog/snapshots.json", "快照时间轴和快照 ID", "是"],
        ["realtime_catalog/observations.parquet", "观测值表", "是"],
        ["realtime_catalog/status.json", "发布状态和最新快照指针", "是，最后写入"],
        ["simulated/simulated_observations_v1.csv", "最小字段样例", "否"],
        ["simulated/simulated_api_fixture_v1.json", "演示分区与预测样例", "否"],
        ["manifest.json", "版本、边界和文件清单", "否"],
    ], [5.6, 7.0, 3.4], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])

    add_heading(doc, "1.1 数据身份与使用边界", 2)
    add_bullet(doc, "实时目录的文件是比赛随源码提供的离线运行样本，不等同于企业生产数据，也不代表永久有效的站点数量。")
    add_bullet(doc, "模拟样例带有 simulation_only 边界，只用于页面联调和演示，不能用于真实监管、告警发布或模型效果结论。")
    add_bullet(doc, "系统不会在实时目录不可用时静默切换到模拟数据；应先看 /api/v1/realtime/status 的 available 和 freshness_status。")

    add_heading(doc, "2 离线数据接入方式")
    add_body(doc, "推荐比赛评审使用环境变量指定数据目录。这样不需要把数据复制到源码内部，也便于企业后续替换为自己的发布目录。")
    add_code(doc, [
        "$env:TAIHU_REALTIME_CATALOG_DIR = (Resolve-Path '.\\企业提交材料\\A23_小型运行数据包_V1.0\\realtime_catalog')",
        "backend\\.venv\\Scripts\\python.exe -m uvicorn backend.main:app --host 127.0.0.1 --port 8000",
        "# 检查：http://127.0.0.1:8000/api/v1/realtime/status",
    ])
    add_body(doc, "如果不设置环境变量，后端默认读取 data-cleaning/storage/silver/mee_realtime。企业部署时可以将同一套四文件放到该默认目录，也可以继续使用环境变量切换目录。")

    add_heading(doc, "2.1 发布顺序", 2)
    add_number = lambda n, t: add_bullet(doc, f"步骤 {n}：{t}")
    add_number(1, "准备 stations.json、snapshots.json 和 observations.parquet。")
    add_number(2, "校验 observations.parquet 的字段、时间格式和 snapshot_id 引用关系。")
    add_number(3, "确认 snapshots.json 最后一条 snapshot_id 与 status.json 的 latest_snapshot_id 相同。")
    add_number(4, "前三个文件写入并校验成功后，最后替换 status.json；这样前端不会读到半发布目录。")

    doc.add_page_break()
    add_heading(doc, "3 字段契约")
    add_heading(doc, "3.1 observations.parquet 最小字段", 2)
    add_table(doc, ["字段", "类型", "说明"], [
        ["spatial_entity_id", "string", "站点或空间对象 ID，必须能在 stations.json 找到"],
        ["observed_at", "datetime/string", "上游观测时间，建议 ISO 8601 并带时区"],
        ["variable_code", "string", "指标代码，例如 chlorophyll_a、total_phosphorus"],
        ["value", "number/null", "观测值，缺测使用 null，不用 0 代替"],
        ["unit", "string", "规范单位，例如 mg/L、degC、ug/L"],
        ["observation_status", "string", "ok、missing、qc_rejected、parse_failed 等"],
        ["qc_status", "string", "质量控制状态"],
        ["retrieved_at", "datetime/string", "本系统抓取或接收时间"],
    ], [4.6, 3.0, 8.4], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_body(doc, "兼容旧样例时，clean_value、value_origin、quality_status 和 is_imputed 可以保留；当前实时 Provider 以发布目录中的 observations.parquet 为准，企业接入前应以运行时 schema 和接口返回为最终校验依据。")

    add_heading(doc, "3.2 stations.json 和 status.json 要点", 2)
    add_bullet(doc, "stations.json 顶层应包含 stations 数组；每个对象至少有可稳定引用的站点 ID、名称、经纬度和坐标可信状态。")
    add_bullet(doc, "snapshots.json 顶层应包含 snapshots 数组；每条记录至少有 snapshot_id 和 observed_at 或时间范围信息。")
    add_bullet(doc, "status.json 必须包含 collection_status、latest_snapshot_id、latest_observed_at、last_success_at 等状态字段。")
    add_bullet(doc, "实时目录缺失、快照不一致或观测文件不可读时，接口应显式返回不可用或错误，不得伪造实时数据。")

    add_heading(doc, "4 接入后的检查")
    add_table(doc, ["检查", "命令或接口", "通过标准"], [
        ["目录检查", "Get-ChildItem realtime_catalog", "四个必需文件均存在"],
        ["链路状态", "GET /api/v1/realtime/status", "HTTP 200，available 和 freshness_status 可解释"],
        ["站点读取", "GET /api/v1/spatial-entities?mode=observed", "返回站点目录或明确错误"],
        ["观测读取", "GET /api/v1/spatial-entities/{entity_id}/observations", "时间、指标、单位和质量字段一致"],
        ["实时汇总", "GET /api/v1/realtime/summary", "数据身份与 as_of 可追溯"],
        ["前端检查", "浏览器刷新页面", "不把 simulated 显示为 observed，不把错误静默成空数据"],
    ], [4.0, 7.0, 5.0], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "5 外部实时接口接入说明")
    add_body(doc, "当前前端不会直接调用企业外部数据源。外部接口由采集与清洗程序负责，采集成功后生成上述发布目录，再由后端读取。企业只需要提供稳定的上游接口、字段映射、鉴权方式和采集频率；不要让浏览器直接暴露上游密钥。")
    add_code(doc, [
        "上游接口 -> 采集适配器 -> 字段映射与质量检查",
        "            -> stations.json / snapshots.json / observations.parquet",
        "            -> status.json（最后发布）",
        "            -> FastAPI /api/v1/realtime/* -> 前端页面",
    ])
    add_body(doc, "比赛版只要求评审能够使用随源码提供的离线数据包跑通页面和接口。若后续接入企业实时数据，应同步更新数据版本、字段字典、采集失败策略和接口文档。")

    add_heading(doc, "6 版本与责任边界")
    add_body(doc, "本数据包与当前后端 OpenAPI 2.1.0、接口文档 V1.0 配套。只要修改文件名、字段、版本或发布规则，就应同步更新 manifest、数据接入说明和接口文档。正式算法接入后，预测接口还应增加模型版本、输入质量门禁和可复现实验记录。")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
